from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

from Error import Error


IMPORT_SHEET = "Variables"
SETUP_SHEET = "Setup"


def create_template(path: str | Path) -> Path:
    workbook = Workbook()
    setup_sheet = workbook.active
    setup_sheet.title = SETUP_SHEET
    setup_sheet["A1"] = "Formula"
    setup_sheet["B1"] = "U / I"
    setup_sheet["A2"] = "Result name"
    setup_sheet["B2"] = "R"
    setup_sheet["A3"] = "Comment"
    setup_sheet["B3"] = "Редактируйте формулу и таблицу измерений на втором листе."
    setup_sheet.column_dimensions["A"].width = 18
    setup_sheet.column_dimensions["B"].width = 45

    data_sheet = workbook.create_sheet(IMPORT_SHEET)
    headers = ["Variable", "Instrumental Error", "Unit", "Measurement 1", "Measurement 2", "Measurement 3"]
    for column_index, value in enumerate(headers, start=1):
        cell = data_sheet.cell(row=1, column=column_index, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        data_sheet.column_dimensions[cell.column_letter].width = 18

    sample_rows = [
        ["U", 0.05, "V", 12.1, 12.2, 12.0],
        ["I", 0.002, "A", 0.251, 0.249, 0.252],
    ]
    for row_index, row_data in enumerate(sample_rows, start=2):
        for column_index, value in enumerate(row_data, start=1):
            data_sheet.cell(row=row_index, column=column_index, value=value)

    output_path = Path(path)
    workbook.save(output_path)
    return output_path


def import_measurements(path: str | Path) -> dict:
    workbook = load_workbook(path, data_only=True)
    if SETUP_SHEET not in workbook.sheetnames or IMPORT_SHEET not in workbook.sheetnames:
        raise ValueError("В Excel-файле должны быть листы 'Setup' и 'Variables'.")

    formula = str(workbook[SETUP_SHEET]["B1"].value or "").strip()
    result_name = str(workbook[SETUP_SHEET]["B2"].value or "Result").strip()
    if not formula:
        raise ValueError("В ячейке Setup!B1 должна быть указана формула.")

    variables_sheet = workbook[IMPORT_SHEET]
    variables: List[dict] = []

    for row in variables_sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        name = str(row[0]).strip()
        instrumental_error = float(row[1] or 0.0)
        unit = str(row[2] or "").strip()
        measurements = [float(value) for value in row[3:] if value is not None and str(value).strip() != ""]
        if not measurements:
            raise ValueError(f"Для переменной '{name}' нет измерений.")

        variables.append(
            {
                "name": name,
                "instrumental_error": instrumental_error,
                "unit": unit,
                "measurements": measurements,
            }
        )

    if not variables:
        raise ValueError("Не найдено ни одной переменной на листе Variables.")

    return {"formula": formula, "result_name": result_name, "variables": variables}


def export_report(
    path: str | Path,
    formula: str,
    result_name: str,
    variables: Iterable[dict],
    direct_measurements: Dict[str, Error],
    indirect_results: Dict[str, Dict],
    graph_images: Dict[str, bytes] | None = None,
) -> Path:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet["A1"] = "Result"
    summary_sheet["B1"] = result_name
    summary_sheet["A2"] = "Formula"
    summary_sheet["B2"] = formula
    summary_sheet["A1"].font = Font(bold=True)
    summary_sheet["A2"].font = Font(bold=True)
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 32

    variables_sheet = workbook.create_sheet("Direct measurements")
    _write_headers(
        variables_sheet,
        [
            "Variable",
            "Unit",
            "Count",
            "Mean",
            "Std",
            "Confidence interval",
            "Instrumental error",
            "Absolute error",
            "Relative error, %",
            "Measurements",
        ],
    )

    row_index = 2
    for variable in variables:
        name = variable["name"]
        error_obj = direct_measurements[name]
        summary = error_obj.summary()
        values = [
            name,
            variable.get("unit", ""),
            summary["count"],
            summary["mean"],
            summary["std"],
            summary["confidence_interval"],
            summary["instrumental_error"],
            summary["absolute_error"],
            summary["relative_error"],
            ", ".join(format(measurement, ".6g") for measurement in variable["measurements"]),
        ]
        for column_index, value in enumerate(values, start=1):
            variables_sheet.cell(row=row_index, column=column_index, value=value)
        row_index += 1

    methods_sheet = workbook.create_sheet("Indirect methods")
    _write_headers(methods_sheet, ["Method", "Value", "Absolute error", "Relative error, %", "Details"])
    row_index = 2
    for result in indirect_results.values():
        values = [
            result["method"],
            result["indirect_value"],
            result["absolute_error"],
            result["relative_error"],
            result["result_string"],
        ]
        for column_index, value in enumerate(values, start=1):
            methods_sheet.cell(row=row_index, column=column_index, value=value)
        row_index += 1

    if graph_images:
        graphs_sheet = workbook.create_sheet("Graphs")
        graphs_sheet["A1"] = "Exported graphs"
        graphs_sheet["A1"].font = Font(bold=True)
        row_cursor = 3
        for title, image_bytes in graph_images.items():
            graphs_sheet.cell(row=row_cursor, column=1, value=title)
            image = XLImage(BytesIO(image_bytes))
            image.width = 720
            image.height = 420
            graphs_sheet.add_image(image, f"A{row_cursor + 1}")
            row_cursor += 24

    output_path = Path(path)
    workbook.save(output_path)
    return output_path


def _write_headers(sheet, headers: List[str]) -> None:
    for column_index, value in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[cell.column_letter].width = max(14, len(value) + 2)
